import os, json
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

rows=list(csv.DictReader(open('/tmp/outputs/defaulters_raw.csv')))
sh={r['code']:r for r in csv.DictReader(open('/tmp/outputs/shifts_raw.csv'),delimiter='|')}
HRS={'12-Hour Shift':('08:00','20:00'),'10-Hour Shift':('09:00','19:00'),'9-Hour Shift':('09:30','18:30'),'Housekeeping 7:30- 4:30':('07:30','16:30'),'Day Shift':('08:00','20:00')}
def m(t): h,mi=t.split(':'); return int(h)*60+int(mi)
for r in rows:
    for k in ['ctr','swd','sns','sl','slc','slm','se','sem','tot','awd','alc','ae','atot']: r[k]=int(r[k])
    r['grp']='Contract' if (r['ctr']==1 or r['typ']=='Contract') else 'Company'
    s=sh.get(r['code'],{})
    r['ms']=s.get('master_shift',''); r['mh']=s.get('master_hours','')
    used=s.get('main_detected',''); base=used.replace(' (night)','')
    r['used']=(used+f" — {HRS[base][0]}-{HRS[base][1]}") if base in HRS else used
    r['mix']=s.get('mixes',''); r['ai']=s.get('avg_in',''); r['ao']=s.get('avg_out','')
    chk=[]
    if r['ms'] and not r['ms'].startswith('(none') and base and base!=r['ms'] and base!='Day Shift':
        chk.append(f"Judged on {base} ({HRS[base][0]}-{HRS[base][1]}) but master says {r['ms']} ({r['mh']}) — lates/earlies are mis-measured")
    if base=='Day Shift' and r['ms'] and r['ms']!='Day Shift' and not r['ms'].startswith('(none'):
        chk.append(f"System used 'Day Shift' though master is {r['ms']}")
    if r['ms'].startswith('(none'): chk.append('No shift set in employee master')
    nights=sum(int(p.rsplit(' ',1)[1]) for p in r['mix'].split(' / ') if '(night)' in p and p.strip())
    tot=sum(int(p.rsplit(' ',1)[1]) for p in r['mix'].split(' / ') if p.strip())
    night_note=f' (mostly night shift: {nights}/{tot} days)' if tot and nights/tot>=0.5 else ''
    if base in HRS and r['ai'] and r['ao']:
        st,en=m(HRS[base][0]),m(HRS[base][1]); ai,ao=m(r['ai']),m(r['ao'])
        if ai<=st+5 and ao<en-45 and (ao-ai)>=(en-st)-90 and (en-st)-(ao-ai)<=90 and ai<st-20:
            chk.append(f'Works an earlier pattern (~{r["ai"]}–{r["ao"]}) than the shift — shift setup likely wrong')
        elif ai<st-20 and ao<en-30:
            chk.append(f'Comes and leaves earlier than shift (~{r["ai"]}–{r["ao"]}) — check shift')
    r['chk']=('; '.join(chk)+night_note) if chk else ('OK — shift setup consistent'+night_note)

SCH={}
for t in open('sched.txt').read().strip().split(','):
    c,m_,h=t.split(':'); SCH[c]=(int(m_),float(h))
EMP={}
for t in open('emps.txt').read().strip().split('|'):
    c,d,ty,ct,de=t.split(':',4); EMP[c]=dict(dept=d,typ=ty,ctr=int(ct),des=de.strip())
isload=lambda c: any(k in EMP.get(c,{}).get('des','').upper() for k in ('LOAD','LODING'))
EXC={l.split(',')[0]:(int(l.split(',')[1]),int(l.split(',')[2]),int(l.split(',')[3])) for l in open('excused.txt').read().split()}
for r in rows:
    raw,ex,exm=EXC.get(r['code'],(r['sl'],0,0))
    r['exc']=ex; r['excm']=exm
    ex=max(ex,r['sl']-r['slc']); r['exc']=ex
    r['slc']=r['sl']-ex   # chargeable = late days minus mornings after a late evening (previous worked day)
    r['slm_ch']=r['slm']-exm
for r in rows:
    r['des']=EMP.get(r['code'],{}).get('des','')
    r['load']=isload(r['code'])
    r['sched'],r['shh']=SCH.get(r['code'],(0,12.0))
    r['lmc']=0 if r['load'] else r['slm_ch']
    r['lcc']=0 if r['load'] else r['slc']
    r['tot2']=r['lmc']+r['sem']
ARTEFACT=json.load(open('artefacts.json')) if os.path.exists('artefacts.json') else {}  # {code: reason} — private, from config
def pattern(r):
    p=[]
    if r['code'] in ARTEFACT: return ARTEFACT[r['code']]
    Lc,E=r['lcc'],r['se']
    if r['load']: p.append('Loading staff — late coming not assessed')
    if r['sl'] and not r['load'] and r['slm']/r['sl']>240: p.append('Very long lates (avg >4 h) — check shift / half-day')
    if r['sns'] and E/r['sns']>=0.8 and Lc<4: p.append('Early every day, rarely late')
    elif r['sns'] and E/r['sns']>=0.8 and r['swd'] and r['sl']/r['swd']>=0.8: p.append('Late AND early almost every day')
    if Lc>=10: p.append('Habitual late (≥10 chargeable)')
    if r['swd']<5: p.append('Few worked days — small sample')
    return '; '.join(p)
def cat(r):
    if r['lcc']>=4 and r['se']>=4: return 'Both'
    if r['lcc']>=4: return 'Late'
    if r['se']>=4: return 'Early'
    return 'Minor'

main=sorted([r for r in rows if r['code'] not in ARTEFACT and r['tot2']>0],key=lambda r:-(r['tot2']/(r['shh']*60)))
F='Arial'
hdr_font=Font(name=F,bold=True,color='FFFFFF'); hdr_fill=PatternFill('solid',fgColor='1F3864')
base_f=Font(name=F,size=10); bold=Font(name=F,size=10,bold=True)
brd=Border(bottom=Side(style='thin',color='D9D9D9'))
fills={'Both':'F8CBAD','Late':'FFE699','Early':'BDD7EE','Minor':'FFFFFF'}

cols=['Rank','Code','Name','Department','Designation','Group','Loading staff (late not assessed)','Shift in employee master','Shift system used in Sep (hours)','Sep shift mix (days)','Avg IN (day shifts)','Avg OUT (day shifts)','Shift check',
'Sep worked days','Avg shift length (h)','Scheduled workday minutes','Late days (≥10 min)','Excused: stayed late previous evening','Chargeable lates (counted)','Late minutes (counted)','Early exits (>15 min, Mon–Sat)','Early minutes','Total minutes lost','% of workday time lost','Workdays lost','Late + early days','Category','Deduction days @ 4 lates = 1 day','Aug chargeable lates','Aug early exits','Aug total minutes','Change in minutes vs Aug','Pattern / flag']
C={n:L(i+1) for i,n in enumerate(cols)}
wb=Workbook(); ws=wb.active; ws.title='Defaulters Sep 2026'
ws['A1']='Late coming & early exit defaulters — September 2026 (ranked by workdays lost; loading staff late-exempt)'; ws['A1'].font=Font(name=F,bold=True,size=13)
ws['A2']='Source: HR SQL Console (attendance_processed, employees, shifts), pulled 10-Oct-2026. Worked days only, miss-punches excluded. See Notes tab.'; ws['A2'].font=Font(name=F,italic=True,size=9,color='595959')
R0=4
for i,h in enumerate(cols,1):
    c=ws.cell(R0,i,h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
ws.row_dimensions[R0].height=48
warn=PatternFill('solid',fgColor='F4B084')
for n,r in enumerate(main,1):
    rr=R0+n
    v={'Rank':n,'Code':int(r['code']),'Name':r['name'].title(),'Department':r['dept'].title(),'Group':r['grp'],
       'Shift in employee master':(f"{r['ms']} ({r['mh']})" if r['mh'] else r['ms']),'Shift system used in Sep (hours)':r['used'],'Sep shift mix (days)':r['mix'],
       'Avg IN (day shifts)':r['ai'] or '—','Avg OUT (day shifts)':r['ao'] or '—','Shift check':r['chk'],
       'Designation':r['des'].title(),'Loading staff (late not assessed)':'Yes' if r['load'] else '',
       'Sep worked days':r['swd'],'Avg shift length (h)':r['shh'],'Scheduled workday minutes':r['sched'],'Late days (≥10 min)':r['sl'],'Excused: stayed late previous evening':r['exc'],
       'Chargeable lates (counted)':('exempt' if r['load'] else r['slc']),'Late minutes (counted)':r['lmc'],'Early exits (>15 min, Mon–Sat)':r['se'],'Early minutes':r['sem'],
       'Total minutes lost':f"={C['Late minutes (counted)']}{rr}+{C['Early minutes']}{rr}",
       '% of workday time lost':f"=IF({C['Scheduled workday minutes']}{rr}=0,0,{C['Total minutes lost']}{rr}/{C['Scheduled workday minutes']}{rr})",
       'Workdays lost':f"=ROUND({C['Total minutes lost']}{rr}/({C['Avg shift length (h)']}{rr}*60),2)",
       'Late + early days':f"=N({C['Chargeable lates (counted)']}{rr})+{C['Early exits (>15 min, Mon–Sat)']}{rr}",
       'Category':cat(r),
       'Deduction days @ 4 lates = 1 day':(('n/a (loading)' if r['load'] else f"=IF({C['Group']}{rr}=\"Company\",FLOOR({C['Chargeable lates (counted)']}{rr}/4,0.25),\"n/a (contractor)\")")),
       'Aug chargeable lates':r['alc'] if r['awd'] else None,'Aug early exits':r['ae'] if r['awd'] else None,'Aug total minutes':r['atot'] if r['awd'] else None,
       'Change in minutes vs Aug':f"=IF({C['Aug total minutes']}{rr}=\"\",\"new in Sep\",{C['Total minutes lost']}{rr}-{C['Aug total minutes']}{rr})",
       'Pattern / flag':pattern(r)}
    for i,name in enumerate(cols,1):
        c=ws.cell(rr,i,v[name]); c.font=base_f; c.border=brd
        if name not in ('Name','Department','Shift in employee master','Shift system used in Sep (hours)','Sep shift mix (days)','Shift check','Pattern / flag'): c.alignment=Alignment(horizontal='center')
        if name in ('Late minutes (counted)','Early minutes','Total minutes lost','Aug total minutes','Scheduled workday minutes'): c.number_format='#,##0'
        if name=='% of workday time lost': c.number_format='0.0%'
        if name=='Workdays lost': c.number_format='0.00'
        if name=='Change in minutes vs Aug': c.number_format='+#,##0;-#,##0;0'
    ws.cell(rr,cols.index('Category')+1).fill=PatternFill('solid',fgColor=fills[cat(r)])
    if not r['chk'].startswith('OK'): ws.cell(rr,cols.index('Shift check')+1).fill=warn
    if r['lcc']>=10 or r['tot2']>=1000: ws.cell(rr,3).font=bold
last=R0+len(main); tr=last+1
ws.cell(tr,3,'TOTAL').font=bold
for name in ['Sep worked days','Avg shift length (h)','Scheduled workday minutes','Late days (≥10 min)','Excused: stayed late previous evening','Chargeable lates (counted)','Late minutes (counted)','Early exits (>15 min, Mon–Sat)','Early minutes','Total minutes lost','% of workday time lost','Workdays lost','Late + early days','Aug chargeable lates','Aug early exits','Aug total minutes']:
    c=ws.cell(tr,cols.index(name)+1,f"=SUM({C[name]}{R0+1}:{C[name]}{last})"); c.font=bold; c.number_format='#,##0'
c=ws.cell(tr,cols.index('% of workday time lost')+1,f"={C['Total minutes lost']}{tr}/{C['Scheduled workday minutes']}{tr}"); c.font=bold; c.number_format='0.0%'
ws.cell(tr,cols.index('Workdays lost')+1).number_format='#,##0.0'
ws.freeze_panes=ws.cell(R0+1,4)
ws.auto_filter.ref=f"A{R0}:{L(len(cols))}{last}"
W={'Designation':18,'Loading staff (late not assessed)':9,'Rank':6,'Code':8,'Name':22,'Department':15,'Group':10,'Shift in employee master':26,'Shift system used in Sep (hours)':28,'Sep shift mix (days)':30,'Avg IN (day shifts)':9,'Avg OUT (day shifts)':9,'Shift check':55,'Pattern / flag':45}
for name in cols: ws.column_dimensions[C[name]].width=W.get(name,10)

# Shift summary sheet
s2=wb.create_sheet('Shift issues')
s2['A1']='Defaulters whose shift setup needs checking before any deduction'; s2['A1'].font=Font(name=F,bold=True,size=13)
hd=['Rank','Code','Name','Department','Shift in master','Shift system used','Avg IN','Avg OUT','Chargeable lates','Early exits','Shift check']
for i,h in enumerate(hd,1):
    c=s2.cell(3,i,h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,horizontal='center')
k=4
for n,r in enumerate(main,1):
    if r['chk'].startswith('OK'): continue
    for i,v in enumerate([n,int(r['code']),r['name'].title(),r['dept'].title(),f"{r['ms']} ({r['mh']})" if r['mh'] else r['ms'],r['used'],r['ai'] or '—',r['ao'] or '—',('exempt' if r['load'] else r['slc']),r['se'],r['chk']],1):
        c=s2.cell(k,i,v); c.font=base_f; c.border=brd
    k+=1
for ac,why in ARTEFACT.items():
    m=[x for x in rows if x['code']==ac]
    if not m: continue
    r=m[0]
    for i,v in enumerate(['excl.',int(ac),r['name'].title(),r['dept'].title(),f"{r['ms']} ({r['mh']})",r['used'],r['ai'],r['ao'],r['slc'],r['se'],why],1):
        c=s2.cell(k,i,v); c.font=base_f
    k+=1
for col,w in zip('ABCDEFGHIJK',[6,8,22,15,28,30,8,8,9,9,80]): s2.column_dimensions[col].width=w
s2.freeze_panes='D4'

# All staff (population incl. zero-default people)
al=wb.create_sheet('All staff Sep')
ah=['Code','Department','Designation','Group','Loading staff','Scheduled workday minutes','Avg shift length (h)','Late minutes (counted)','Early minutes','Total minutes lost','Defaulter (any late/early)','Category','Workdays lost']
for i,h in enumerate(ah,1):
    c=al.cell(1,i,h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,horizontal='center')
byc={r['code']:r for r in rows}
k=2
for code,(sm,shh) in sorted(SCH.items()):
    e=EMP.get(code,{}); r=byc.get(code)
    grp='Contract' if (e.get('ctr')==1 or e.get('typ')=='Contract') else 'Company'
    lm=(r['lmc'] if r else 0); em=(r['sem'] if r else 0)
    if code in ARTEFACT: lm=0; em=0
    vals=[int(code),e.get('dept','?').title(),e.get('des','').title(),grp,'Yes' if isload(code) else '',sm,shh,lm,em,f'=H{k}+I{k}','Yes' if (lm+em)>0 else '',cat(r) if (r and code not in ARTEFACT and r['tot2']>0) else '',f'=IF(G{k}=0,0,J{k}/(G{k}*60))']
    for i,v in enumerate(vals,1):
        c=al.cell(k,i,v); c.font=base_f
    k+=1
alast=k-1
for col,w in zip('ABCDEFGHIJKLM',[8,18,22,10,8,12,8,10,10,10,10,10,10]): al.column_dimensions[col].width=w
for i in range(2,alast+1): al[f'M{i}'].number_format='0.00'
al.freeze_panes='B2'; al.auto_filter.ref=f'A1:M{alast}'

dp=wb.create_sheet('By Department',1)
dp['A1']='Workday time lost to late coming + early exit, by department — September 2026'; dp['A1'].font=Font(name=F,bold=True,size=13)
dp['A2']='Denominator = scheduled shift minutes of EVERYONE in the department who worked in Sep (not just defaulters). Loading staff: late not counted.'; dp['A2'].font=Font(name=F,italic=True,size=9,color='595959')
dh=['Department','Group','Staff who worked','Staff with any late/early','Scheduled workday minutes','Late minutes (counted)','Early minutes','Total minutes lost','% of workday time lost','Workdays lost','Both late & early (≥4 each)','Share of staff with any late/early']
for i,h in enumerate(dh,1):
    c=dp.cell(4,i,h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
dp.row_dimensions[4].height=42
A=lambda col:f"'All staff Sep'!${col}$2:${col}${alast}"
depts={}
for code in SCH:
    e=EMP.get(code,{}); grp='Contract' if (e.get('ctr')==1 or e.get('typ')=='Contract') else 'Company'
    depts.setdefault(e.get('dept','?').title(),set()).add(grp)
# order by minutes lost
def dlost(d): return sum((byc[c]['tot2'] if c in byc and c not in ARTEFACT else 0) for c in SCH if EMP.get(c,{}).get('dept','?').title()==d)
k=5
for d in sorted(depts,key=lambda d:-dlost(d)):
    g='/'.join(sorted(depts[d]))
    vals=[d,g,f'=COUNTIF({A("B")},A{k})',f'=COUNTIFS({A("B")},A{k},{A("K")},"Yes")',f'=SUMIF({A("B")},A{k},{A("F")})',f'=SUMIF({A("B")},A{k},{A("H")})',f'=SUMIF({A("B")},A{k},{A("I")})',f'=F{k}+G{k}',f'=IF(E{k}=0,0,H{k}/E{k})',f'=ROUND(SUMIF({A("B")},A{k},{A("M")}),1)',f'=COUNTIFS({A("B")},A{k},{A("L")},"Both")',f'=IF(C{k}=0,0,D{k}/C{k})']
    for i,v in enumerate(vals,1):
        c=dp.cell(k,i,v); c.font=base_f; c.border=brd
        if i>=3: c.alignment=Alignment(horizontal='center')
    for col in 'EFGH': dp[f'{col}{k}'].number_format='#,##0'
    dp[f'I{k}'].number_format='0.0%'; dp[f'L{k}'].number_format='0%'
    k+=1
dept_tot=k
dp.cell(k,1,'TOTAL').font=bold
for col in 'CDEFGHJK':
    c=dp[f'{col}{k}']; c.value=f'=SUM({col}5:{col}{k-1})'; c.font=bold; c.number_format='#,##0.0' if col=='J' else '#,##0'
dp[f'I{k}']=f'=H{k}/E{k}'; dp[f'I{k}'].number_format='0.0%'; dp[f'I{k}'].font=bold
dp[f'L{k}']=f'=D{k}/C{k}'; dp[f'L{k}'].number_format='0%'; dp[f'L{k}'].font=bold
for col,w in zip('ABCDEFGHIJKL',[20,16,9,10,13,11,10,11,10,12,9,12]): dp.column_dimensions[col].width=w
dp.freeze_panes='B5'

# Summary
s=wb.create_sheet('Summary')
s['A1']='Summary — September 2026 defaulters'; s['A1'].font=Font(name=F,bold=True,size=13)
D="'Defaulters Sep 2026'"; rng=lambda name:f"{D}!${C[name]}${R0+1}:${C[name]}${last}"
s.append([]); s.append(['Category','Company','Contract','All','Minutes lost (all)'])
for c in s[3]: c.font=hdr_font; c.fill=hdr_fill
for key,label in [('Both','Both late & early (≥4 each)'),('Late','Late only (≥4 chargeable)'),('Early','Early only (≥4)'),('Minor','Minor (<4 of each)')]:
    rr=s.max_row+1
    s.append([label,f'=COUNTIFS({rng("Category")},"{key}",{rng("Group")},"Company")',f'=COUNTIFS({rng("Category")},"{key}",{rng("Group")},"Contract")',f'=B{rr}+C{rr}',f'=SUMIFS({rng("Total minutes lost")},{rng("Category")},"{key}")'])
rr=s.max_row+1
s.append(['Total defaulters',f'=SUM(B4:B{rr-1})',f'=SUM(C4:C{rr-1})',f'=SUM(D4:D{rr-1})',f'=SUM(E4:E{rr-1})'])
s.append([])
s.append(['Total hours lost (all defaulters)',f'=ROUND(E{rr}/60,0)'])
s.append(['Workdays lost (each person at own shift length)',f'=ROUND(SUM({rng("Workdays lost")}),1)'])
s.append(['Scheduled workday minutes, ALL staff who worked in Sep',f"='By Department'!E{dept_tot}"])
s.append(['% of all scheduled workday time lost to late + early',f"='By Department'!I{dept_tot}"])
s.append(['Habitual late (≥10 chargeable lates, excl. loading)',f'=COUNTIF({rng("Chargeable lates (counted)")},">=10")'])
s.append(['Company staff eligible for late deduction (≥4 chargeable, excl. loading)',f'=COUNTIFS({rng("Chargeable lates (counted)")},">=4",{rng("Group")},"Company")'])
s.append(['Indicative deduction days for Company staff (4 lates = 1 day)',f'=SUMIFS({rng("Deduction days @ 4 lates = 1 day")},{rng("Group")},"Company")'])
s.append(['Defaulters with a shift issue to check first',f'=COUNTIF({rng("Shift check")},"<>OK*")'])
s.append(['   …of which measured against a different shift than the master',f'=COUNTIF({rng("Shift check")},"*mis-measured*")'])
s.append(['   …of which keep an earlier timing than their shift',f'=COUNTIF({rng("Shift check")},"*earlier*")'])
for row in s.iter_rows(min_row=4):
    for c in row:
        if isinstance(c.value,str) and c.value.startswith("='By Department'!I"): c.number_format='0.0%'
        if isinstance(c.value,str) and c.value.startswith("='By Department'!E"): c.number_format='#,##0'
for row in s.iter_rows(min_row=4):
    for c in row:
        if not c.font.bold: c.font=base_f
for c in s[rr]: c.font=bold
s.column_dimensions['A'].width=66
for col in 'BCDE': s.column_dimensions[col].width=16

nt=wb.create_sheet('Notes')
notes=['Definitions','Period: 1–30 Sep 2026 (Aug 2026 shown for comparison). Only days marked P, WOP, ½P or WO½P; miss-punch days excluded.',
'Late day: flagged late by ≥10 minutes (grace is 9 min).','Chargeable late: excludes a morning after the employee left ≥20 min past shift end on their previous WORKED day (rule locked 04-Sep-2026; previous worked day used so a late Saturday covers Monday). Excused mornings are removed from the late count AND late minutes, hours lost, % and workdays lost.',
'Early exit: >15 minutes early, Monday–Saturday; readings of 600+ minutes (night-shift misreads) excluded.',
'Ranking: workdays lost = (counted late minutes + early-exit minutes) ÷ the person’s own shift length. % of workday time lost = minutes lost ÷ scheduled shift minutes for the days worked (½-days count half).','Loading staff (designation contains LOADING/LODING): late coming is NOT assessed — their late minutes are excluded; early exits still count.',
'Shift in employee master: the default shift on the employee record. Shift system used: the shift the attendance engine actually matched on most September days — late/early is measured against THIS one.',
'Sep shift mix: how many September worked days were matched to each shift (night = overnight pairing).',
'Avg IN / Avg OUT: average punch times on day-shift days only (night days excluded, as their times cross midnight).',
'Shift check: "mis-measured" = system used a different shift than the master; "earlier pattern" = person keeps a consistent earlier timing that fits a full shift, so the shift setup is the likely problem, not discipline.',
'Deduction days are indicative (4 lates = 1 day, quarter-days); contractors are not deducted in payroll.',
'','Excluded from ranking']+[f'{c}: {w}' for c,w in ARTEFACT.items()]
for i,t in enumerate(notes,1):
    c=nt.cell(i,1,t); c.font=Font(name=F,bold=True,size=11) if t in ('Definitions','Excluded from ranking') else base_f; c.alignment=Alignment(wrap_text=True)
nt.column_dimensions['A'].width=130
wb.save('/tmp/outputs/Late_Early_Defaulters_Sep2026.xlsx')
print(len(main), sum(1 for r in main if not r['chk'].startswith('OK')))
