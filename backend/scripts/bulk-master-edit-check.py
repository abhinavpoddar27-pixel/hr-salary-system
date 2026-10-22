#!/usr/bin/env python3
"""P5-02 — employee master bulk edit: browser check (Chromium, built dist, real hr / finance / admin logins).

Scratch DATA_DIR, fictional people only (T8001–T8008, "Test Person A–H"). Flow:
  hr ticks 3 people on Employee Master → "Set field for 3" → shift 10HR → Save → audit rows (same as bulk-shift) →
  Sheet: download department + bank columns → edit (openpyxl) → upload → stale row + bad account flagged →
  leave out → "Save 1 now · Send 1 to finance" → undo the direct batch → leaving for a loan borrower (sent) →
  a hand-made sheet without _meta shows the warning → 390 px →
  finance opens Bulk batches, sees the full account, approves → four bank columns written, masked audit →
  approves the leaving → Left + loan flagged → admin's own batch refused (SELF_APPROVAL). 0 page errors.
Usage: python3 backend/scripts/bulk-master-edit-check.py [--base ROOT_OF_BASE_WORKTREE] [screenshot_dir]
  --base: another checkout's build: no "Bulk edit" button (old "Assign Shift" window), no master-edit kind.
Port: PORT env (default 3506).
"""
import json, os, re, sqlite3, subprocess, sys, tempfile, time, urllib.request, urllib.error
import openpyxl
from playwright.sync_api import sync_playwright

args = sys.argv[1:]
BASE_ROOT = None
if '--base' in args:
    i = args.index('--base'); BASE_ROOT = os.path.abspath(args[i + 1]); del args[i:i + 2]
SHOTS = args[0] if args else None
ROOT = BASE_ROOT or os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
PORT = int(os.environ.get('PORT', '3506')); BASE = f'http://127.0.0.1:{PORT}'
WORK = tempfile.mkdtemp(prefix='bulk-me-'); DB = os.path.join(WORK, 'hr_system.db')
PW = {'admin': 'Admin@123', 'hr': 'Indriyan@2025', 'finance': 'Finance@2025'}
C = 'Indriyan Beverages Pvt Ltd'
PASS = FAIL = 0


def check(label, expected, actual):
    global PASS, FAIL
    if expected == actual: PASS += 1; print(f'  ✓ {label}')
    else: FAIL += 1; print(f'  ✗ {label} — expected {expected!r}, got {actual!r}')


def q(sql, *a):
    con = sqlite3.connect(DB); r = con.execute(sql, a).fetchall(); con.close(); return r


def ex(sql, *a):
    con = sqlite3.connect(DB); con.execute(sql, a); con.commit(); con.close()


def http(method, path, tok=None, body=None):
    data = json.dumps(body).encode() if body is not None else None
    h = {'Content-Type': 'application/json'}
    if tok: h['Authorization'] = f'Bearer {tok}'
    req = urllib.request.Request(BASE + path, data=data, method=method, headers=h)
    try:
        with urllib.request.urlopen(req, timeout=60) as r: return r.status, json.loads(r.read() or b'{}')
    except urllib.error.HTTPError as e:
        try: return e.code, json.loads(e.read() or b'{}')
        except Exception: return e.code, {}


def login_api(u):
    s, b = http('POST', '/api/auth/login', body={'username': u, 'password': PW[u]})
    return b.get('token')


PEOPLE = [('T8001', 'A'), ('T8002', 'B'), ('T8003', 'C'), ('T8004', 'D'), ('T8005', 'E'), ('T8006', 'F'), ('T8007', 'G'), ('T8008', 'H')]


def seed():
    c = sqlite3.connect(DB)
    sid = c.execute("SELECT id FROM shifts WHERE code='12HR'").fetchone()[0]
    for code, letter in PEOPLE:
        c.execute("INSERT INTO employees(code,name,department,designation,company,status,employment_type,gross_salary,date_of_joining,"
                  "default_shift_id,shift_code,account_number,bank_account,ifsc_code,ifsc) VALUES(?,?,'TEST DEPT','TEST OPERATOR',?,'Active','Permanent',"
                  "20000,'2024-01-01',?,'12HR','1111000011','1111000011','SBIN0000001','SBIN0000001')", (code, f'Test Person {letter}', C, sid))
    # T8005: an active loan → leaving flags it for recovery.
    lid = c.execute("""INSERT INTO loans (borrower_type, employee_code, company, loan_type, principal_amount, tenure_months, emi_amount,
        status, requested_by, remaining_balance) VALUES ('plant', 'T8005', ?, 'Personal', 6000, 3, 2000, 'active', 'hr', 6000)""", (C,)).lastrowid
    for i in range(1, 4):
        c.execute('INSERT INTO loan_instalments (loan_id, sequence, due_month, due_year, amount_due) VALUES (?, ?, ?, 2026, 2000)', (lid, i, 9 + i))
    c.commit(); c.close()


env = dict(os.environ, DATA_DIR=WORK, JWT_SECRET='p502-secret', PORT=str(PORT), NODE_ENV='production',
           ADMIN_PASSWORD=PW['admin'], HR_PASSWORD=PW['hr'], FINANCE_PASSWORD=PW['finance'])
srv = None


def start():
    global srv
    srv = subprocess.Popen(['node', f'{ROOT}/backend/server.js'], env=env, stdout=open(f'{WORK}/log', 'a'), stderr=subprocess.STDOUT)
    for _ in range(80):
        try: urllib.request.urlopen(BASE + '/api/version', timeout=2); return
        except Exception: time.sleep(.5)
    raise SystemExit('server did not start')


def stop():
    srv.terminate(); srv.wait(timeout=10)


def emp(code, cols):
    return q(f'SELECT {cols} FROM employees WHERE code=?', code)[0]


try:
    start(); seed()
    with sync_playwright() as p:
        b = p.chromium.launch()
        errs, cerrs, bad = [], [], []

        def login(user, width=1440):
            ctx = b.new_context(viewport={'width': width, 'height': 900}, accept_downloads=True)
            pg = ctx.new_page()
            pg.on('pageerror', lambda e: errs.append(str(e)))
            pg.on('console', lambda m: cerrs.append(m.text) if m.type == 'error' else None)
            pg.on('response', lambda r: bad.append(f'{r.status} {r.request.method} {r.url.split(BASE)[-1]}') if '/api/' in r.url and r.status >= 400 else None)
            pg.goto(BASE + '/login'); pg.get_by_placeholder('admin').fill(user); pg.get_by_placeholder('••••••••').fill(PW[user])
            pg.get_by_role('button', name='Sign in').click(); pg.wait_for_url(lambda u: '/login' not in u)
            return pg

        def tick(pg, code):
            pg.locator('tr', has_text=code).first.locator('input[type=checkbox]').check()

        def wait_dialog(pg):
            pg.wait_for_selector('[role=dialog]'); time.sleep(.45)  # useConfirm ignores clicks in its first 300 ms
            return pg.get_by_role('dialog')

        if BASE_ROOT:
            print('\n— BASE build —')
            hr = login('hr'); hr.goto(BASE + '/employees'); hr.wait_for_load_state('networkidle'); time.sleep(.8)
            check('base: no "Bulk edit" button on Employee Master', 0, hr.locator('[data-testid=me-open]').count())
            tick(hr, 'T8001')
            check('base: the old "Assign Shift" window button is there', 1, hr.get_by_role('button', name='Assign Shift').count())
            s, _ = http('POST', '/api/bulk/master-edit/preview', login_api('admin'), {'common': {}, 'rows': [{'code': 'T8001', 'changes': {'designation': 'X'}}]})
            check('base: no master-edit kind (404)', 404, s)
            b.close(); raise SystemExit(0)

        print('\n— hr: tick 3 people → Set field for 3 → shift 10HR → Save —')
        hr = login('hr')
        hr.goto(BASE + '/employees'); hr.wait_for_load_state('networkidle'); time.sleep(.6)
        check('old "Assign Shift" window button is gone', 0, hr.get_by_role('button', name='Assign Shift').count())
        for code in ('T8001', 'T8002', 'T8003'): tick(hr, code)
        hr.locator('[data-testid=me-set-field]').click()
        hr.wait_for_selector('[data-testid=me-panel]')
        check('panel opens inline (no dialog)', 0, hr.get_by_role('dialog').count())
        hr.wait_for_function("document.querySelector('[data-testid=me-field]') && document.querySelector('[data-testid=me-field]').value === 'shift'")
        opts = hr.locator('[data-testid=me-value] option').all_inner_texts()
        check('shift list hides the legacy codes (as the old window did)', False, any(o.startswith('DAY ') or o.startswith('GEN ') for o in opts))
        hr.locator('[data-testid=me-value]').select_option('10HR')
        hr.wait_for_selector('[data-testid=me-counts-direct]'); time.sleep(1)
        check('3 ready in "Saved now"', True, hr.locator('[data-testid=me-counts-direct]').inner_text().startswith('3 ready'))
        check('row shows Shift: 12HR → 10HR', True, '12HR' in hr.locator('tr[data-code=T8001]').inner_text() and '10HR' in hr.locator('tr[data-code=T8001]').inner_text())
        save_btn = hr.locator('[data-testid=me-save]')
        check('save label', 'Save 3 now', save_btn.inner_text().strip())
        save_btn.dblclick()
        wait_dialog(hr).get_by_role('button', name='Save').click()
        hr.wait_for_selector('[data-testid=me-result]'); time.sleep(.5)
        check('one direct batch applied (double click → 1)', [(1, 'applied')], q("SELECT COUNT(*), MAX(status) FROM bulk_batches WHERE kind='master-edit'"))
        check('shift written for the 3', [('10HR',)] * 3, q("SELECT shift_code FROM employees WHERE code IN ('T8001','T8002','T8003') ORDER BY code"))
        check('3 shift_assignment audit rows, same text as bulk-shift', [(3, 'Bulk shift change from 12HR to 10HR', 'shift_change', 'employee_master')],
              q("SELECT COUNT(*), MAX(remark), MAX(action_type), MAX(stage) FROM audit_log WHERE field_name='shift_assignment'"))
        hr.locator('[data-testid=me-result]').get_by_role('button', name='Close').click(); time.sleep(.3)

        print('\n— hr: sheet → download department + bank → edit → upload —')
        hr.locator('[data-testid=me-open]').click(); hr.wait_for_selector('[data-testid=me-panel]')
        hr.locator('[data-testid=me-mode-sheet]').click()
        hr.locator('[data-field=shift]').uncheck()
        for f in ('department', 'account', 'ifsc'): hr.locator(f'[data-field={f}]').check()
        with hr.expect_download() as dl:
            hr.locator('[data-testid=me-download]').click()
        path = os.path.join(WORK, 'sheet.xlsx'); dl.value.save_as(path)
        wb = openpyxl.load_workbook(path)
        check('workbook has Employees, Shifts and a hidden _meta sheet', (['Employees', 'Shifts', '_meta'], 'hidden'), (wb.sheetnames, wb['_meta'].sheet_state))
        ws = wb['Employees']
        hdr = [c.value for c in ws[1]]
        check('columns: Code, Name, Department, Bank account, IFSC', ['Code', 'Name', 'Department', 'Bank account', 'IFSC'], hdr)
        rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
        check('8 people exported with current values', (8, '1111000011', 'TEST DEPT'), (len(rows), ws.cell(rows['T8002'], 4).value, ws.cell(rows['T8001'], 3).value))
        ws.cell(rows['T8001'], 3).value = 'TEST STORES'           # direct, new department → warning
        ws.cell(rows['T8002'], 4).value = '2222000022'            # finance
        ws.cell(rows['T8002'], 5).value = 'hdfc0000002'
        ws.cell(rows['T8003'], 4).value = '99AB'                  # bad account → error
        ws.cell(rows['T8004'], 3).value = 'TEST PACKING'          # stale: master moves below
        wb.save(path)
        ex("UPDATE employees SET department='MOVED BY SOMEONE' WHERE code='T8004'")
        hr.locator('[data-testid=me-file]').set_input_files(path)
        hr.wait_for_selector('[data-testid=me-sheet-line]'); time.sleep(1.5)
        check('4 changed people read (unchanged cells ignored)', True, '4 people with changes' in hr.locator('[data-testid=me-sheet-line]').inner_text())
        check('stale row T8004 refused (changed since download)', True, 'changed since you downloaded' in hr.locator('[data-testid=me-section-direct] tr[data-code=T8004]').inner_text())
        check('bad account T8003 refused', True, '6 to 20 digits' in hr.locator('[data-testid=me-section-details] tr[data-code=T8003]').inner_text())
        check('Save disabled while problems are there', True, hr.locator('[data-testid=me-save]').is_disabled())
        hr.locator('[data-testid=me-leave-out-direct]').click(); time.sleep(1.2)
        hr.locator('[data-testid=me-leave-out-details]').click(); time.sleep(1.4)
        if SHOTS: hr.screenshot(path=os.path.join(SHOTS, 'p502-hr-sheet.png'), full_page=True)
        check('save label after leaving out', 'Save 1 now · Send 1 to finance', hr.locator('[data-testid=me-save]').inner_text().strip())
        hr.locator('[data-testid=me-save]').click()
        wait_dialog(hr).get_by_role('button', name='Save').click()
        hr.wait_for_selector('[data-testid=me-result]'); time.sleep(.5)
        check('direct department written, finance one not yet', ('TEST STORES', '1111000011'), (emp('T8001', 'department')[0], emp('T8002', 'account_number')[0]))
        check('a details batch waits for finance', [(1,)], q("SELECT COUNT(*) FROM bulk_batches WHERE kind='master-edit-sensitive' AND status='submitted'"))
        hr.locator('[data-testid=me-result]').get_by_role('button', name='Undo').click()
        dlg = wait_dialog(hr)
        dlg.locator('input, textarea').first.fill('wrong department')
        dlg.get_by_role('button', name='Undo batch').click(); time.sleep(1.2)
        check('undo puts the department back', 'TEST DEPT', emp('T8001', 'department')[0])
        hr.locator('[data-testid=me-result]').get_by_role('button', name='Close').click(); time.sleep(.3)

        print('\n— hr: leaving for a loan borrower → sent to finance —')
        hr.goto(BASE + '/employees'); hr.wait_for_load_state('networkidle'); time.sleep(.5)
        tick(hr, 'T8005'); hr.locator('[data-testid=me-set-field]').click(); hr.wait_for_selector('[data-testid=me-panel]')
        hr.locator('[data-testid=me-field]').select_option('exit')
        hr.locator('[data-testid=me-exit-date]').fill('2026-10-05'); hr.locator('[data-testid=me-exit-reason]').fill('Resigned (test)')
        hr.wait_for_selector('[data-testid=me-section-exit] tr[data-code=T8005]'); time.sleep(1.2)
        check('leaving row warns about the open loan', True, 'open loan(s) will be flagged' in hr.locator('[data-testid=me-section-exit] tr[data-code=T8005]').inner_text())
        hr.locator('[data-testid=me-save]').click()
        wait_dialog(hr).get_by_role('button', name='Send to finance').click()
        hr.wait_for_selector('[data-testid=me-result]'); time.sleep(.4)
        check('T8005 still Active until finance approves', 'Active', emp('T8005', 'status')[0])
        hr.locator('[data-testid=me-result]').get_by_role('button', name='Close').click(); time.sleep(.3)

        print('\n— hr: a hand-made sheet without _meta —')
        hm = openpyxl.Workbook(); hs = hm.active; hs.append(['Code', 'Designation', 'Gross']); hs.append(['T8006', 'TEST PACKER', '']); hs.append(['T8007', 'TEST OPERATOR', '25000'])
        hpath = os.path.join(WORK, 'hand.xlsx'); hm.save(hpath)
        hr.locator('[data-testid=me-open]').click(); hr.wait_for_selector('[data-testid=me-panel]')
        hr.locator('[data-testid=me-mode-sheet]').click()
        hr.locator('[data-testid=me-file]').set_input_files(hpath); time.sleep(1.6)
        check('no-_meta warning shown', 1, hr.locator('[data-testid=me-nometa]').count())
        check('Gross column refused with the reason', True, 'Bulk salary revision' in hr.locator('[data-testid=me-section-direct] tr[data-code=T8007]').inner_text())
        hr.locator('[data-testid=me-panel]').get_by_role('button', name='Close').first.click(); time.sleep(.3)

        print('\n— hr at 390 px —')
        hr390 = login('hr', width=390)
        hr390.goto(BASE + '/employees'); hr390.wait_for_load_state('networkidle'); hr390.locator('[data-testid=me-open]').click()
        hr390.wait_for_selector('[data-testid=me-panel]')
        right = hr390.locator('[data-testid=me-panel]').evaluate('e => e.getBoundingClientRect().right')
        check('390: panel fits the screen width', True, right <= 391)
        if SHOTS: hr390.screenshot(path=os.path.join(SHOTS, 'p502-390.png'), full_page=True)

        print('\n— finance: Bulk batches → full account → approve both —')
        stop(); start()  # login limiter
        fin = login('finance')
        fin.goto(BASE + '/finance-verification?tab=bulk'); fin.wait_for_load_state('networkidle'); time.sleep(.8)
        D = q("SELECT id FROM bulk_batches WHERE kind='master-edit-sensitive' AND json_extract(common_json,'$.section')='details'")[0][0]
        E = q("SELECT id FROM bulk_batches WHERE kind='master-edit-sensitive' AND json_extract(common_json,'$.section')='exit'")[0][0]
        row = fin.locator(f'tr[data-batch-id="{D}"]')
        row.get_by_role('button', name='Open').click(); fin.wait_for_selector('[data-testid=bulk-queue-rows]'); time.sleep(.3)
        rt = fin.locator('[data-testid=bulk-queue-rows]').inner_text()
        check('finance sees the full account and IFSC', True, '1111000011 → 2222000022' in rt and 'HDFC0000002' in rt)
        if SHOTS: fin.screenshot(path=os.path.join(SHOTS, 'p502-finance.png'), full_page=True)
        for bid in (D, E):
            fin.locator(f'tr[data-batch-id="{bid}"]').locator('input[type=checkbox]').check()
        fin.locator('[data-testid=bulk-bar]').get_by_role('button', name=re.compile('Approve 2')).click(); time.sleep(.5)
        fin.get_by_role('dialog').get_by_role('button', name='Approve and apply').click(); time.sleep(2)
        check('both batches applied by finance', [('applied', 'finance'), ('applied', 'finance')], q(f'SELECT status, approved_by FROM bulk_batches WHERE id IN ({D},{E}) ORDER BY id'))
        check('bank: all four columns', ('2222000022', '2222000022', 'HDFC0000002', 'HDFC0000002'), emp('T8002', 'account_number, bank_account, ifsc_code, ifsc'))
        check('bank audit masked to last 4', [('****0011', '****0022')], q("SELECT old_value, new_value FROM audit_log WHERE stage='bank_details_change' AND field_name='account_number'"))
        check('T8005 Left on 2026-10-05', ('Left', '2026-10-05'), emp('T8005', 'status, date_of_exit'))
        check('T8005 loan flagged recover_at_exit', [('recover_at_exit', 1)], q("SELECT status, exit_flag FROM loans WHERE employee_code='T8005'"))

        print('\n— admin cannot approve their own batch —')
        atok = login_api('admin')
        s, pv = http('POST', '/api/bulk/master-edit-sensitive/preview', atok, {'common': {}, 'rows': [{'code': 'T8008', 'changes': {'bank_name': 'TEST BANK'}}]})
        g = pv['groups'][0]
        s, sub = http('POST', '/api/bulk/master-edit-sensitive/submit', atok, {'common': {}, 'planHash': g['planHash'], 'rows': [pv['rows'][0]['input']]})
        s, body = http('POST', f"/api/bulk/batches/{sub.get('batchId')}/approve", atok, {})
        check('403 SELF_APPROVAL', (403, 'SELF_APPROVAL'), (s, body.get('code')))
        check('bank name not written', None, emp('T8008', 'bank_name')[0])

        check('drift 0 (net = gross_earned − total_deductions)', [(0,)], q('SELECT COUNT(*) FROM salary_computations WHERE ABS(net_salary-(gross_earned-total_deductions))>1'))
        print('\n— Errors —')
        check('0 page errors', [], errs)
        check('0 API ≥ 400 (browser)', [], bad)
        check('0 console errors', [], [c for c in cerrs if 'favicon' not in c and 'vite.svg' not in c])
        b.close()
except SystemExit:
    raise
except Exception:
    import traceback; traceback.print_exc(); FAIL += 1
finally:
    try: stop()
    except Exception: pass
    print(f'\n{PASS} passed, {FAIL} failed  (scratch {WORK})')
    sys.exit(1 if FAIL else 0)
